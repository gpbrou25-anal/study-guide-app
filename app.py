"""
Study Guide — turns uploaded lecture files (PDF / PPTX / DOCX) into an
interactive study guide: structured summary, plain-language explanations,
a glossary of hard terms, and a quiz. Remembers your progress and quiz
scores across sessions.

Storage: uses a Postgres database (e.g. free Neon) when a DATABASE_URL
secret/env var is set — the right choice when deployed on Streamlit
Community Cloud, since that filesystem can reset. Falls back to a local
SQLite file automatically when no DATABASE_URL is configured, so it still
works with zero setup when you just run it on your own machine.

Run with:  streamlit run app.py
Needs an Anthropic API key (env var ANTHROPIC_API_KEY, or paste it in the sidebar).
"""

import io
import json
import os
import time
from datetime import datetime

import streamlit as st
from sqlalchemy import create_engine, text

# --- optional heavy imports guarded so the app still boots with a clear error ---
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None
try:
    from pptx import Presentation
except ImportError:
    Presentation = None
try:
    from docx import Document
except ImportError:
    Document = None
try:
    import anthropic
except ImportError:
    anthropic = None

LOCAL_DB_PATH = os.path.join(os.path.dirname(__file__), "study_guide.db")
MAX_CHARS = 400_000  # safety cap on how much extracted text we send per course

MODEL_OPTIONS = {
    "Sonnet 5 (recommended — best balance)": "claude-sonnet-5",
    "Opus 5 (most capable, slower/pricier)": "claude-opus-5",
    "Haiku 4.5 (fastest, cheapest)": "claude-haiku-4-5-20251001",
}

# ---------------------------------------------------------------------------
# Database — Postgres (e.g. free Neon) when DATABASE_URL is configured,
# local SQLite file otherwise. Same schema and queries work on both.
# ---------------------------------------------------------------------------

def _database_url() -> str:
    # st.secrets raises if no secrets.toml exists at all — guard that.
    try:
        url = st.secrets.get("DATABASE_URL")
    except Exception:
        url = None
    url = url or os.environ.get("DATABASE_URL")
    if url:
        # SQLAlchemy wants the psycopg driver spelled out; Neon/most hosts give
        # a plain "postgres://" or "postgresql://" string.
        if url.startswith("postgres://"):
            url = "postgresql+psycopg2://" + url[len("postgres://"):]
        elif url.startswith("postgresql://"):
            url = "postgresql+psycopg2://" + url[len("postgresql://"):]
        return url
    return f"sqlite:///{LOCAL_DB_PATH}"


@st.cache_resource
def get_engine():
    engine = create_engine(_database_url(), pool_pre_ping=True)
    is_pg = engine.dialect.name == "postgresql"
    id_col = "id SERIAL PRIMARY KEY" if is_pg else "id INTEGER PRIMARY KEY AUTOINCREMENT"
    with engine.begin() as conn:
        conn.execute(text(f"""CREATE TABLE IF NOT EXISTS courses(
            {id_col},
            name TEXT NOT NULL,
            created_at TEXT NOT NULL,
            content_json TEXT NOT NULL
        )"""))
        conn.execute(text("""CREATE TABLE IF NOT EXISTS progress(
            course_id INTEGER NOT NULL,
            section_index INTEGER NOT NULL,
            studied INTEGER NOT NULL DEFAULT 0,
            PRIMARY KEY (course_id, section_index)
        )"""))
        conn.execute(text(f"""CREATE TABLE IF NOT EXISTS quiz_attempts(
            {id_col},
            course_id INTEGER NOT NULL,
            timestamp TEXT NOT NULL,
            score INTEGER NOT NULL,
            total INTEGER NOT NULL
        )"""))
    return engine


def save_course(engine, name, content: dict) -> int:
    with engine.begin() as conn:
        result = conn.execute(
            text("INSERT INTO courses(name, created_at, content_json) VALUES (:n, :c, :j) RETURNING id"),
            {"n": name, "c": datetime.now().isoformat(timespec="seconds"), "j": json.dumps(content)},
        )
        return result.scalar_one()


def load_courses(engine):
    with engine.connect() as conn:
        rows = conn.execute(text("SELECT id, name, created_at FROM courses ORDER BY created_at DESC")).fetchall()
    return rows


def load_course_content(engine, course_id) -> dict:
    with engine.connect() as conn:
        row = conn.execute(text("SELECT content_json FROM courses WHERE id=:id"), {"id": course_id}).fetchone()
    return json.loads(row[0]) if row else None


def delete_course(engine, course_id):
    with engine.begin() as conn:
        conn.execute(text("DELETE FROM courses WHERE id=:id"), {"id": course_id})
        conn.execute(text("DELETE FROM progress WHERE course_id=:id"), {"id": course_id})
        conn.execute(text("DELETE FROM quiz_attempts WHERE course_id=:id"), {"id": course_id})


def set_studied(engine, course_id, section_index, studied: bool):
    with engine.begin() as conn:
        conn.execute(
            text(
                "INSERT INTO progress(course_id, section_index, studied) VALUES (:c, :s, :v) "
                "ON CONFLICT(course_id, section_index) DO UPDATE SET studied=excluded.studied"
            ),
            {"c": course_id, "s": section_index, "v": int(studied)},
        )


def get_studied_set(engine, course_id):
    with engine.connect() as conn:
        rows = conn.execute(
            text("SELECT section_index FROM progress WHERE course_id=:id AND studied=1"), {"id": course_id}
        ).fetchall()
    return {r[0] for r in rows}


def save_quiz_attempt(engine, course_id, score, total):
    with engine.begin() as conn:
        conn.execute(
            text("INSERT INTO quiz_attempts(course_id, timestamp, score, total) VALUES (:c, :t, :s, :tot)"),
            {"c": course_id, "t": datetime.now().isoformat(timespec="seconds"), "s": score, "tot": total},
        )


def get_quiz_history(engine, course_id):
    with engine.connect() as conn:
        return conn.execute(
            text("SELECT timestamp, score, total FROM quiz_attempts WHERE course_id=:id ORDER BY timestamp DESC"),
            {"id": course_id},
        ).fetchall()


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.read()
    if name.endswith(".pdf"):
        if PdfReader is None:
            raise RuntimeError("pypdf is not installed — run: pip install pypdf")
        reader = PdfReader(io.BytesIO(data))
        parts = []
        for i, page in enumerate(reader.pages):
            txt = page.extract_text() or ""
            if txt.strip():
                parts.append(f"[Page {i + 1}]\n{txt}")
        return "\n\n".join(parts)
    elif name.endswith(".pptx"):
        if Presentation is None:
            raise RuntimeError("python-pptx is not installed — run: pip install python-pptx")
        prs = Presentation(io.BytesIO(data))
        parts = []
        for i, slide in enumerate(prs.slides):
            slide_txt = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    slide_txt.append(shape.text_frame.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        slide_txt.append(" | ".join(c.text for c in row.cells))
            # speaker notes often carry the actual explanation — include them
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
                if notes.strip():
                    slide_txt.append(f"(Speaker notes: {notes})")
            if slide_txt:
                parts.append(f"[Slide {i + 1}]\n" + "\n".join(slide_txt))
        return "\n\n".join(parts)
    elif name.endswith(".docx"):
        if Document is None:
            raise RuntimeError("python-docx is not installed — run: pip install python-docx")
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    else:
        raise RuntimeError(f"Unsupported file type: {uploaded_file.name}")


# ---------------------------------------------------------------------------
# Claude API call — turns raw extracted text into the structured study guide
# ---------------------------------------------------------------------------

SYSTEM_PROMPT = """You are an expert study-guide writer. You turn raw, messy lecture \
material (extracted from slides/PDFs/Word docs, so formatting is imperfect) into a \
complete, structured, genuinely useful study guide for a student preparing for an exam.

Rules:
- Cover EVERYTHING in the source material. Do not skip topics, slides, or examples \
because they seem minor — the student wants full coverage, not a highlights reel.
- Organize into logical sections/chapters that follow the material's own structure \
(don't invent a different structure than what's actually there).
- For each section, write a real explanation in plain, clear language — as if \
teaching someone who missed the lecture, not repeating slide bullet points verbatim.
- Any jargon, acronym, or technical term must be explained in the section's glossary.
- Preserve and explain every concrete example, case study, formula, or numeric \
figure found in the material — these are exactly what exams test.
- Write in the same language as the source material (if the slides are in French, \
write the guide in French; if in English, write in English; etc.).
- Generate quiz questions covering the ENTIRE material, roughly 1-3 questions per \
section depending on its density, all multiple-choice with exactly 4 options, one \
correct answer, and a one-sentence explanation of why the correct answer is right.

Respond with ONLY a single JSON object, no other text, no markdown fences, matching \
exactly this shape:
{
  "course_name": "string",
  "language": "string (e.g. 'English', 'Français', 'Español')",
  "sections": [
    {
      "title": "string",
      "summary": "a thorough, multi-paragraph explanation in plain language (use \\n\\n between paragraphs)",
      "key_points": ["string", "..."],
      "examples": [{"title": "string", "explanation": "string"}],
      "glossary": [{"term": "string", "definition": "string"}]
    }
  ],
  "quiz": [
    {
      "question": "string",
      "options": ["string", "string", "string", "string"],
      "correct_index": 0,
      "explanation": "string"
    }
  ]
}"""


def build_study_guide(api_key, model, course_name, raw_text):
    if anthropic is None:
        raise RuntimeError("The 'anthropic' package is not installed — run: pip install anthropic")
    if len(raw_text) > MAX_CHARS:
        raw_text = raw_text[:MAX_CHARS]
        st.warning(
            f"The uploaded material was very long and was truncated to the first "
            f"{MAX_CHARS:,} characters to stay within a safe request size."
        )
    client = anthropic.Anthropic(api_key=api_key)
    user_prompt = (
        f"Course name: {course_name}\n\n"
        f"Raw extracted lecture material follows. Build the complete study guide JSON "
        f"as instructed.\n\n---\n{raw_text}\n---"
    )
    message = client.messages.create(
        model=model,
        max_tokens=16000,
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": user_prompt}],
    )
    text = "".join(block.text for block in message.content if block.type == "text")
    text = text.strip()
    # tolerate the model wrapping in a code fence despite instructions
    if text.startswith("```"):
        text = text.split("```")[1]
        if text.startswith("json"):
            text = text[4:]
    return json.loads(text.strip())


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------

st.set_page_config(page_title="Study Guide", page_icon="\U0001F4DA", layout="wide")

conn = get_engine()

with st.sidebar:
    st.title("\U0001F4DA Study Guide")

    st.subheader("API key")
    try:
        default_key = st.secrets.get("ANTHROPIC_API_KEY", "")
    except Exception:
        default_key = ""
    default_key = default_key or os.environ.get("ANTHROPIC_API_KEY", "")
    api_key = st.text_input(
        "Anthropic API key",
        value=st.session_state.get("api_key", default_key),
        type="password",
        help="Not saved to disk — only kept for this session. Set ANTHROPIC_API_KEY as an "
             "environment variable to skip pasting it every time.",
    )
    st.session_state["api_key"] = api_key

    model_label = st.selectbox("Model", list(MODEL_OPTIONS.keys()))
    model = MODEL_OPTIONS[model_label]

    st.divider()
    st.subheader("Add a course")
    new_course_name = st.text_input("Course name")
    uploaded_files = st.file_uploader(
        "Upload lecture files (PDF, PPTX, DOCX — as many as you have for this course)",
        type=["pdf", "pptx", "docx"],
        accept_multiple_files=True,
    )
    process_clicked = st.button("Build study guide", type="primary", use_container_width=True)

    st.divider()
    st.subheader("Your courses")
    courses = load_courses(conn)
    course_labels = {c[0]: f"{c[1]}  ({c[2][:10]})" for c in courses}
    selected_course_id = None
    if courses:
        selected_course_id = st.radio(
            "Pick a course to study",
            options=[c[0] for c in courses],
            format_func=lambda cid: course_labels[cid],
            label_visibility="collapsed",
        )
        if st.button("Delete selected course", use_container_width=True):
            delete_course(conn, selected_course_id)
            st.rerun()
    else:
        st.caption("No courses yet — upload your first one above.")

# --- process upload ---
if process_clicked:
    if not api_key:
        st.error("Paste your Anthropic API key in the sidebar first.")
    elif not new_course_name:
        st.error("Give the course a name.")
    elif not uploaded_files:
        st.error("Upload at least one file.")
    else:
        with st.spinner(f"Reading {len(uploaded_files)} file(s) and building your study guide — this can take a minute or two for a full course..."):
            try:
                combined_text = []
                for f in uploaded_files:
                    combined_text.append(f"\n\n===== FILE: {f.name} =====\n\n" + extract_text(f))
                raw_text = "".join(combined_text)
                if not raw_text.strip():
                    st.error("Couldn't extract any text from those files (they may be scanned images).")
                else:
                    content = build_study_guide(api_key, model, new_course_name, raw_text)
                    course_id = save_course(conn, new_course_name, content)
                    st.success(f"'{new_course_name}' is ready — {len(content.get('sections', []))} sections, {len(content.get('quiz', []))} quiz questions.")
                    st.session_state["just_built"] = course_id
                    time.sleep(0.5)
                    st.rerun()
            except json.JSONDecodeError:
                st.error("The AI's response wasn't valid JSON — this can happen on very dense material. Try again, or with fewer files at once.")
            except Exception as e:
                st.error(f"Something went wrong: {e}")

# --- main area ---
if not courses:
    st.title("Welcome")
    st.write(
        "Upload your first course's lecture files in the sidebar (PDF, PPTX or DOCX — "
        "you can add several files at once for one course) and click **Build study guide**. "
        "Everything gets extracted, explained, and turned into a quiz automatically."
    )
else:
    course_id = st.session_state.get("just_built", selected_course_id)
    content = load_course_content(conn, course_id)
    st.title(content.get("course_name", "Course"))

    studied = get_studied_set(conn, course_id)
    n_sections = len(content.get("sections", []))
    st.progress(len(studied) / n_sections if n_sections else 0, text=f"{len(studied)}/{n_sections} sections studied")

    tab_guide, tab_glossary, tab_quiz, tab_progress = st.tabs(
        ["\U0001F4D6 Study guide", "\U0001F520 Glossary", "\U0001F9E0 Quiz", "\U0001F4CA Progress"]
    )

    with tab_guide:
        search = st.text_input("Search within this course", key="guide_search")
        for i, sec in enumerate(content.get("sections", [])):
            haystack = (sec.get("title", "") + sec.get("summary", "")).lower()
            if search and search.lower() not in haystack:
                continue
            is_studied = i in studied
            label = ("✅ " if is_studied else "⬜ ") + sec.get("title", f"Section {i+1}")
            with st.expander(label):
                st.markdown(sec.get("summary", "").replace("\n", "\n\n"))
                if sec.get("key_points"):
                    st.markdown("**Key points**")
                    for kp in sec["key_points"]:
                        st.markdown(f"- {kp}")
                if sec.get("examples"):
                    st.markdown("**Examples**")
                    for ex in sec["examples"]:
                        st.markdown(f"*{ex.get('title', '')}* — {ex.get('explanation', '')}")
                if sec.get("glossary"):
                    with st.popover("Hard words in this section"):
                        for g in sec["glossary"]:
                            st.markdown(f"**{g.get('term', '')}** — {g.get('definition', '')}")
                checked = st.checkbox("Mark as studied", value=is_studied, key=f"studied_{course_id}_{i}")
                if checked != is_studied:
                    set_studied(conn, course_id, i, checked)
                    st.rerun()

    with tab_glossary:
        all_terms = []
        for sec in content.get("sections", []):
            for g in sec.get("glossary", []):
                all_terms.append((g.get("term", ""), g.get("definition", ""), sec.get("title", "")))
        all_terms.sort(key=lambda t: t[0].lower())
        gsearch = st.text_input("Search glossary", key="glossary_search")
        for term, definition, section_title in all_terms:
            if gsearch and gsearch.lower() not in term.lower():
                continue
            st.markdown(f"**{term}** &nbsp;·&nbsp; *{section_title}*")
            st.caption(definition)

    with tab_quiz:
        quiz = content.get("quiz", [])
        if not quiz:
            st.info("No quiz questions were generated for this course.")
        else:
            with st.form(f"quiz_form_{course_id}"):
                answers = {}
                for i, q in enumerate(quiz):
                    st.markdown(f"**{i + 1}. {q['question']}**")
                    answers[i] = st.radio(
                        f"q{i}", q["options"], key=f"quiz_{course_id}_{i}", label_visibility="collapsed", index=None
                    )
                submitted = st.form_submit_button("Submit quiz", type="primary")
            if submitted:
                score = 0
                for i, q in enumerate(quiz):
                    correct = q["options"][q["correct_index"]]
                    user_answer = answers[i]
                    is_right = user_answer == correct
                    score += int(is_right)
                    icon = "✅" if is_right else "❌"
                    st.markdown(f"{icon} **{i + 1}.** {q['question']}")
                    if not is_right:
                        st.caption(f"You answered: {user_answer or '(no answer)'} — Correct: {correct}")
                    st.caption(q.get("explanation", ""))
                st.success(f"Score: {score}/{len(quiz)}")
                save_quiz_attempt(conn, course_id, score, len(quiz))

    with tab_progress:
        history = get_quiz_history(conn, course_id)
        if history:
            st.markdown("**Quiz history**")
            for ts, score, total in history:
                st.markdown(f"- {ts} — {score}/{total} ({round(100*score/total)}%)")
        else:
            st.caption("No quiz attempts yet.")
        st.markdown("**Sections studied**")
        for i, sec in enumerate(content.get("sections", [])):
            mark = "✅" if i in studied else "⬜"
            st.markdown(f"{mark} {sec.get('title', f'Section {i+1}')}")
